import streamlit as st
import pandas as pd
import io
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Finanzas - Grupo SERVYRE", layout="wide")

def formato_mx(val):
    if pd.isnull(val): return "$0.00"
    try:
        num = float(val)
        partes = f"{num:,.2f}".split(".")
        return f"${int(partes[0].replace(',', '')):,}.{partes[1] if len(partes) > 1 else '00'}"
    except (ValueError, TypeError):
        return str(val)

def generar_excel_ejecutivo(df_datos, empresa_nombre):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Pagos"
    ws.views.sheetView[0].showGridLines = True

    font_titulo = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    fill_titulo = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    font_empresa = Font(name="Calibri", size=12, bold=True, color="1F497D")
    fill_empresa = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    font_prov = Font(name="Calibri", size=11, bold=True, color="333333")
    fill_prov = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    font_moneda = Font(name="Calibri", size=10, bold=True, color="1F497D")
    fill_moneda = PatternFill(start_color="E9EDF4", end_color="E9EDF4", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    font_normal = Font(name="Calibri", size=10, color="000000")
    font_total = Font(name="Calibri", size=11, bold=True, color="000000")
    fill_total = PatternFill(start_color="8EA9DB", end_color="8EA9DB", fill_type="solid")

    borde_delgado = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
    borde_total = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))

    row_idx = 1
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=8)
    cell = ws.cell(row=row_idx, column=1, value=f"REPORTE EJECUTIVO DE PAGOS — {empresa_nombre.upper()}")
    cell.font = font_titulo; cell.fill = fill_titulo; cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row_idx].height = 35
    row_idx += 2

    headers = ["TIPO", "PROVEEDOR", "FECHA VENCIMIENTO", "FOLIO / DOCUMENTO", "UUID", "MONEDA", "DESCRIPCIÓN", "SALDO PENDIENTE"]
    col_emp = 'EmpresaOrigen' if 'EmpresaOrigen' in df_datos.columns else None
    col_prov = 'BusinessEntityName' if 'BusinessEntityName' in df_datos.columns else None
    col_curr = 'Currency' if 'Currency' in df_datos.columns else None

    empresas = sorted(df_datos[col_emp].unique()) if col_emp else ['General']

    for empresa in empresas:
        df_emp = df_datos[df_datos[col_emp] == empresa] if col_emp else df_datos
        total_empresa = df_emp['Saldo_Pendiente'].sum()

        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
        cell_emp = ws.cell(row=row_idx, column=1, value=f"EMPRESA: {str(empresa).upper()}")
        cell_emp.font = font_empresa; cell_emp.fill = fill_empresa; cell_emp.alignment = Alignment(horizontal="left", vertical="center")
        
        cell_emp_tot = ws.cell(row=row_idx, column=8, value=total_empresa)
        cell_emp_tot.font = font_empresa; cell_emp_tot.fill = fill_empresa; cell_emp_tot.number_format = '"$"#,##0.00'; cell_emp_tot.alignment = Alignment(horizontal="right", vertical="center")
        
        for c_idx in range(1, 9): ws.cell(row=row_idx, column=c_idx).border = borde_delgado
        ws.row_dimensions[row_idx].height = 26
        row_idx += 1

        proveedores = sorted(df_emp[col_prov].unique()) if col_prov else []
        for prov in proveedores:
            df_prov = df_emp[df_emp[col_prov] == prov]
            total_prov = df_prov['Saldo_Pendiente'].sum()

            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
            cell_prov = ws.cell(row=row_idx, column=1, value=f"   👤 {str(prov)}")
            cell_prov.font = font_prov; cell_prov.fill = fill_prov; cell_prov.alignment = Alignment(horizontal="left", vertical="center")
            
            cell_prov_tot = ws.cell(row=row_idx, column=8, value=total_prov)
            cell_prov_tot.font = font_prov; cell_prov_tot.fill = fill_prov; cell_prov_tot.number_format = '"$"#,##0.00'; cell_prov_tot.alignment = Alignment(horizontal="right", vertical="center")
            
            for c_idx in range(1, 9): ws.cell(row=row_idx, column=c_idx).border = borde_delgado
            ws.row_dimensions[row_idx].height = 22
            row_idx += 1

            monedas_prov = sorted(df_prov[col_curr].dropna().unique()) if col_curr else ['MXN']
            for moneda in monedas_prov:
                df_moneda = df_prov[df_prov[col_curr] == moneda] if col_curr else df_prov
                total_moneda = df_moneda['Saldo_Pendiente'].sum()

                ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
                cell_mon = ws.cell(row=row_idx, column=1, value=f"      💱 Moneda: {str(moneda)}")
                cell_mon.font = font_moneda; cell_mon.fill = fill_moneda; cell_mon.alignment = Alignment(horizontal="left", vertical="center")
                
                cell_mon_tot = ws.cell(row=row_idx, column=8, value=total_moneda)
                cell_mon_tot.font = font_moneda; cell_mon_tot.fill = fill_moneda; cell_mon_tot.number_format = '"$"#,##0.00'; cell_mon_tot.alignment = Alignment(horizontal="right", vertical="center")
                
                for c_idx in range(1, 9): ws.cell(row=row_idx, column=c_idx).border = borde_delgado
                ws.row_dimensions[row_idx].height = 20
                row_idx += 1

                for col_num, h in enumerate(headers, 1):
                    c = ws.cell(row=row_idx, column=col_num, value=h)
                    c.font = font_header; c.fill = fill_header; c.alignment = Alignment(horizontal="center", vertical="center"); c.border = borde_delgado
                ws.row_dimensions[row_idx].height = 20
                row_idx += 1

                for _, r in df_moneda.iterrows():
                    ws.cell(row=row_idx, column=1, value=str(r.get('Tipo_Movimiento', ''))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=2, value=str(r.get('BusinessEntityName', ''))).alignment = Alignment(horizontal="left")
                    ws.cell(row=row_idx, column=3, value=str(r.get('Fecha_Fmt', ''))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=4, value=str(r.get('DocFolio', ''))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=5, value=str(r.get('UUID', ''))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=6, value=str(r.get('Currency', 'MXN'))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=7, value=str(r.get('Title', ''))).alignment = Alignment(horizontal="left")
                    
                    c_val = ws.cell(row=row_idx, column=8, value=float(r.get('Saldo_Pendiente', 0)))
                    c_val.number_format = '"$"#,##0.00'; c_val.alignment = Alignment(horizontal="right")

                    for col_num in range(1, 9): ws.cell(row=row_idx, column=col_num).font = font_normal; ws.cell(row=row_idx, column=col_num).border = borde_delgado
                    ws.row_dimensions[row_idx].height = 18
                    row_idx += 1
                row_idx += 1
            row_idx += 1
        row_idx += 1

    ws.cell(row=row_idx, column=1, value="TOTAL GENERAL").font = font_total
    for col_num in range(1, 8): ws.cell(row=row_idx, column=col_num).fill = fill_total; ws.cell(row=row_idx, column=col_num).border = borde_total
    
    cell_tot_gen = ws.cell(row=row_idx, column=8, value=df_datos['Saldo_Pendiente'].sum())
    cell_tot_gen.font = font_total; cell_tot_gen.fill = fill_total; cell_tot_gen.number_format = '"$"#,##0.00'
    cell_tot_gen.border = borde_total; cell_tot_gen.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[row_idx].height = 26

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

ruta_archivo = "Consolidado_Master.xlsx"

@st.cache_data
def cargar_datos_finanzas(path):
    if not os.path.exists(path): return None, None, None, None, None, None
    try:
        xls = pd.ExcelFile(path)
        sheets = xls.sheet_names
        return (
            pd.read_excel(path, sheet_name='FacturaCliente') if 'FacturaCliente' in sheets else pd.DataFrame(),
            pd.read_excel(path, sheet_name='SolicitudPago') if 'SolicitudPago' in sheets else pd.DataFrame(),
            pd.read_excel(path, sheet_name='OrdenCompra') if 'OrdenCompra' in sheets else pd.DataFrame(),
            pd.read_excel(path, sheet_name='EdoCuenta') if 'EdoCuenta' in sheets else pd.DataFrame(),
            pd.read_excel(path, sheet_name='FacturaCompra') if 'FacturaCompra' in sheets else None,
            pd.read_excel(path, sheet_name='Gastos') if 'Gastos' in sheets else None
        )
    except Exception:
        return None, None, None, None, None, None

df_factura, df_tesoreria, df_ordenes, df_edocuenta, df_fact_compra, df_gastos = cargar_datos_finanzas(ruta_archivo)

st.sidebar.title("💰 Módulo de Finanzas")
st.sidebar.markdown("---")
submodulo = st.sidebar.radio("Seleccione Submódulo:", ["Facturación", "OC y SP", "Reporte Pagos"], key="sub_finanzas")

if submodulo == "Facturación":
    st.title("📊 Módulo de Facturación y Cobranza")
    columnas_facturacion = ['EmpresaOrigen', 'BusinessEntityName', 'DocFolio', 'DateDocument', 'Currency', 'SubTotal', 'TotalTax', 'Total', 'TotalPagado', 'SaldoPendiente', 'UUID', 'Status']
    
    if df_factura is not None and not df_factura.empty:
        df_f = df_factura.copy()
        if df_edocuenta is not None and not df_edocuenta.empty and 'DocumentID' in df_f.columns and 'DocumentID' in df_edocuenta.columns:
            cols_e = ['EmpresaOrigen', 'DocumentID', 'Amount'] if 'EmpresaOrigen' in df_edocuenta.columns else ['DocumentID', 'Amount']
            df_e_sub = df_edocuenta[cols_e].copy()
            df_e_sub['Amount'] = pd.to_numeric(df_e_sub['Amount'], errors='coerce').fillna(0)
            if 'EmpresaOrigen' in df_e_sub.columns and 'EmpresaOrigen' in df_f.columns:
                df_pagos = df_e_sub.groupby(['EmpresaOrigen', 'DocumentID'])['Amount'].sum().reset_index()
                df_f = pd.merge(df_f, df_pagos, on=['EmpresaOrigen', 'DocumentID'], how='left')
            else:
                df_pagos = df_e_sub.groupby('DocumentID')['Amount'].sum().reset_index()
                df_f = pd.merge(df_f, df_pagos, on='DocumentID', how='left')
            df_f['TotalPagado'] = df_f['Amount'].fillna(0)
        else:
            df_f['TotalPagado'] = 0.0

        df_f['Total'] = pd.to_numeric(df_f['Total'], errors='coerce').fillna(0)
        df_f['SaldoPendiente'] = (df_f['Total'] - df_f['TotalPagado']).apply(lambda x: max(0.0, x))

        if 'DateDocument' in df_f.columns:
            df_f['DateDocument'] = pd.to_datetime(df_f['DateDocument'], errors='coerce')
            df_f['Año'] = df_f['DateDocument'].dt.year.fillna(0).astype(int)
        else: df_f['Año'] = 0

        st.markdown("#### ⚙️ Filtros de Selección")
        c1, c2, c3 = st.columns(3)
        with c1:
            e_sel = st.multiselect("Filtrar por Empresa Origen:", sorted(df_f['EmpresaOrigen'].dropna().unique()) if 'EmpresaOrigen' in df_f.columns else [], key="f_e")
            if e_sel: df_f = df_f[df_f['EmpresaOrigen'].isin(e_sel)]
        with c2:
            c_sel = st.multiselect("Filtrar por Cliente(s):", sorted(df_f['BusinessEntityName'].dropna().unique()) if 'BusinessEntityName' in df_f.columns else [], key="f_c")
            if c_sel: df_f = df_f[df_f['BusinessEntityName'].isin(c_sel)]
        with c3:
            a_sel = st.multiselect("Filtrar por Año(s):", sorted([int(a) for a in df_f['Año'].unique() if a > 0], reverse=True), key="f_a")
            if a_sel: df_f = df_f[df_f['Año'].isin(a_sel)]

        st.markdown("---")
        st.metric("Total Facturado (Filtrado)", formato_mx(df_f['Total'].sum()))
        
        df_v = df_f.copy()
        for col in ['Total', 'TotalTax', 'SubTotal', 'TotalPagado', 'SaldoPendiente']:
            if col in df_v.columns: df_v[col] = df_v[col].apply(formato_mx)
        st.dataframe(df_v[[c for c in columnas_facturacion if c in df_v.columns]], use_container_width=True)

elif submodulo == "OC y SP":
    st.title("📦 Módulo de Órdenes de Compra y Solicitudes de Pago")
    columnas_oc_sp = ['EmpresaOrigen', 'DocFolio', 'BusinessEntityName', 'DateDocument', 'Title', 'CostCenterName', 'Currency', 'Rate', 'SubTotal', 'Total', 'UUID', 'Saldo_Pendiente']
    
    tab_oc, tab_sp = st.tabs(["OrdenCompra", "SolicitudPago"])
    with tab_oc:
        if df_ordenes is not None and not df_ordenes.empty:
            df_o = df_ordenes.copy()
            if 'DateDocument' in df_o.columns:
                df_o['DateDocument'] = pd.to_datetime(df_o['DateDocument'], errors='coerce')
                df_o['Año'] = df_o['DateDocument'].dt.year.fillna(0).astype(int)
            else: df_o['Año'] = 0

            st.markdown("#### ⚙️ Filtros Orden de Compra")
            c1, c2, c3, c4 = st.columns(4)
            with c1:
                e_sel = st.multiselect("Empresa Origen (OC):", sorted(df_o['EmpresaOrigen'].dropna().unique()) if 'EmpresaOrigen' in df_o.columns else [], key="oc_e")
                if e_sel: df_o = df_o[df_o['EmpresaOrigen'].isin(e_sel)]
            with c2:
                p_sel = st.multiselect("Proveedor (OC):", sorted(df_o['BusinessEntityName'].dropna().unique()) if 'BusinessEntityName' in df_o.columns else [], key="oc_p")
                if p_sel: df_o = df_o[df_o['BusinessEntityName'].isin(p_sel)]
            with c3:
                a_sel = st.multiselect("Año (OC):", sorted([int(a) for a in df_o['Año'].unique() if a > 0], reverse=True), key="oc_a")
                if a_sel: df_o = df_o[df_o['Año'].isin(a_sel)]
            with c4:
                if st.checkbox("Saldo pendiente > $1.00 (OC)", key="oc_s") and 'Saldo_Pendiente' in df_o.columns:
                    df_o = df_o[df_o['Saldo_Pendiente'] > 1.0]

            st.markdown("---")
            st.dataframe(df_o[[c for c in columnas_oc_sp if c in df_o.columns]], use_container_width=True)

    with tab_sp:
        if df_tesoreria is not None and not df_tesoreria.empty:
            df_s = df_tesoreria.copy()
            if 'DateDocument' in df_s.columns:
                df_s['DateDocument'] = pd.to_datetime(df_s['DateDocument'], errors='coerce')
                df_s['Año'] = df_s['DateDocument'].dt.year.fillna(0).astype(int)
            else: df_s['Año'] = 0

            st.markdown("#### ⚙️ Filtros Solicitudes de Pago")
            c1, c2, c3, c4 = st.columns(4)
            with c1:
                e_sel = st.multiselect("Empresa Origen (SP):", sorted(df_s['EmpresaOrigen'].dropna().unique()) if 'EmpresaOrigen' in df_s.columns else [], key="sp_e")
                if e_sel: df_s = df_s[df_s['EmpresaOrigen'].isin(e_sel)]
            with c2:
                p_sel = st.multiselect("Proveedor (SP):", sorted(df_s['BusinessEntityName'].dropna().unique()) if 'BusinessEntityName' in df_s.columns else [], key="sp_p")
                if p_sel: df_s = df_s[df_s['BusinessEntityName'].isin(p_sel)]
            with c3:
                a_sel = st.multiselect("Año (SP):", sorted([int(a) for a in df_s['Año'].unique() if a > 0], reverse=True), key="sp_a")
                if a_sel: df_s = df_s[df_s['Año'].isin(a_sel)]
            with c4:
                if st.checkbox("Saldo pendiente > $1.00 (SP)", key="sp_s") and 'Saldo_Pendiente' in df_s.columns:
                    df_s = df_s[df_s['Saldo_Pendiente'] > 1.0]

            st.markdown("---")
            st.dataframe(df_s[[c for c in columnas_oc_sp if c in df_s.columns]], use_container_width=True)

elif submodulo == "Reporte Pagos":
    st.title("📋 Reporte Ejecutivo de Pagos con Saldo Pendiente")
    
    df_oc_b = df_ordenes.copy() if df_ordenes is not None and not df_ordenes.empty else pd.DataFrame()
    if not df_oc_b.empty: df_oc_b['Tipo_Movimiento'] = 'Orden de Compra (OC)'
    
    df_sp_b = df_tesoreria.copy() if df_tesoreria is not None and not df_tesoreria.empty else pd.DataFrame()
    if not df_sp_b.empty: df_sp_b['Tipo_Movimiento'] = 'Solicitud de Pago (SP)'
    
    df_list = [df for df in [df_oc_b, df_sp_b] if not df.empty]
    if df_list:
        df_tot = pd.concat(df_list, ignore_index=True)
        if 'Saldo_Pendiente' in df_tot.columns:
            df_tot['Saldo_Pendiente'] = pd.to_numeric(df_tot['Saldo_Pendiente'], errors='coerce').fillna(0)
            df_tot = df_tot[df_tot['Saldo_Pendiente'] > 1.0]
        
        st.markdown(f"### 💰 **Total Saldo Pendiente General:** {formato_mx(df_tot['Saldo_Pendiente'].sum())}")
        
        empresa_excel = df_tot['EmpresaOrigen'].iloc[0] if 'EmpresaOrigen' in df_tot.columns and len(df_tot['EmpresaOrigen'].dropna().unique()) == 1 else "Consolidado"
        st.download_button(
            label="📥 Descargar Reporte en Excel con UUID",
            data=generar_excel_ejecutivo(df_tot, empresa_excel),
            file_name="Reporte_Ejecutivo_Pagos_UUID.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        st.markdown("---")
        st.dataframe(df_tot, use_container_width=True)