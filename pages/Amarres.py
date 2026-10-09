import glob
import io
import os
import re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Módulo de Amarres",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.title("🔗 Módulo de Amarres Contables")

# ORDEN PREFERENTE DE EMPRESAS
ORDEN_EMPRESAS_PRIORIDAD = [
    "CIVLAT",
    "SERVYRE",
    "FERVIC",
    "CIV",
    "GPO SERVYRE",
    "FGS",
    "LATIN",
    "EFCO",
    "PESAZA",
    "LIMPIESPIN",
    "SERVYCARGO",
    "CIVMEX",
    "COMANA",
    "PROINA",
    "SEVILAT",
    "IPV",
    "SERSEÑAL",
    "INMOBILIARIA",
    "VIALTECNO",
    "LABORATORIO",
    "FPSB",
    "SIGNAL",
]

MAPEO_CODIGO_EMPRESA = {
    "0468": "CIVLAT",
    "0467": "CIV",
    "0469": "CIVMEX",
    "2872": "SERVYRE",
    "1636": "SERSEÑAL",
    "0813": "EFCO",
    "0942": "FGS",
    "1127": "FERVIC",
    "1216": "FPSB",
    "1144": "PESAZA",
    "1149": "GPO SERVYRE",
    "1404": "INMOBILIARIA",
    "1603": "LABORATORIO",
    "1616": "LATIN",
    "1626": "LIMPIESPIN",
    "2871": "SERVYCARGO",
    "3329": "VIALTECNO",
    "2937": "SEVILAT",
    "2396": "PROINA",
    "1428": "IPV",
    "0665": "COMANA",
}

MAPEO_NOMBRE_A_CODIGO = {v: k for k, v in MAPEO_CODIGO_EMPRESA.items()}


def ordenar_empresas_segun_prioridad(lista_empresas):
    def obtener_posicion(emp_nombre):
        emp_clean = str(emp_nombre).strip().upper()
        for idx, pref in enumerate(ORDEN_EMPRESAS_PRIORIDAD):
            if pref == emp_clean or pref in emp_clean:
                return idx
        return 999

    return sorted(lista_empresas, key=obtener_posicion)


def obtener_estructura_balanzas():
    archivos = glob.glob("Balanzas/**/Balanza.xlsx", recursive=True)
    estructura = []
    for path in archivos:
        path_norm = path.replace("\\", "/")
        partes = path_norm.split("/")
        if len(partes) >= 4:
            anio = partes[-3]
            mes = partes[-2]
            estructura.append({
                'anio': str(anio),
                'mes': str(mes),
                'ruta': path_norm,
                'mtime': os.path.getmtime(path),
            })
    if estructura:
        estructura.sort(key=lambda x: x['mtime'], reverse=True)
    return estructura


@st.cache_data(ttl=600)
def obtener_lista_empresas(ruta):
    xls = pd.ExcelFile(ruta)
    sheets = xls.sheet_names
    if len(sheets) > 1 and "Hoja1" in sheets:
        sheets.remove("Hoja1")
    return ordenar_empresas_segun_prioridad(sheets)


@st.cache_data(ttl=600)
def cargar_hoja_balanza(ruta, nombre_hoja):
    return pd.read_excel(ruta, sheet_name=nombre_hoja)


def parse_monto_robusto(val):
    if pd.isnull(val):
        return 0.0
    val_str = str(val).replace('$', '').replace(',', '').replace(' ', '').strip()
    if not val_str or val_str.lower() in ('nan', 'none', '-'):
        return 0.0
    try:
        return float(val_str)
    except ValueError:
        return 0.0


def extraer_codigo_contraparte_robusto(cta_str):
    segs = cta_str.split('-')
    for seg in reversed(segs):
        seg_clean = seg.strip()
        if seg_clean in MAPEO_CODIGO_EMPRESA:
            return seg_clean
    for seg in reversed(segs):
        seg_clean = seg.strip()
        if seg_clean.isdigit() and len(seg_clean) in (3, 4):
            return seg_clean
    return ""


def exportar_excel_matriz_individual(df_matriz, titulo_reporte, subtitulo_reporte):
    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte"
    ws.views.sheetView[0].showGridLines = True

    TITLE_FONT = Font(name="Calibri", size=14, bold=True, color="1E293B")
    SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="475569")
    HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    REGULAR_FONT = Font(name="Calibri", size=11, color="334155")
    TOTAL_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    TOTAL_FONT = Font(name="Calibri", size=11, bold=True, color="0F172A")

    THIN_BORDER = Border(
        left=Side(style='thin', color='E2E8F0'), right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'), bottom=Side(style='thin', color='E2E8F0')
    )
    TOTAL_BORDER = Border(
        top=Side(style='thin', color='475569'), bottom=Side(style='double', color='0F172A'),
        left=Side(style='thin', color='E2E8F0'), right=Side(style='thin', color='E2E8F0')
    )

    ws.cell(row=1, column=1, value=titulo_reporte).font = TITLE_FONT
    ws.cell(row=2, column=1, value=subtitulo_reporte).font = SUBTITLE_FONT

    start_row = 4
    ws.cell(row=start_row, column=1, value="EMPRESA").fill = HEADER_FILL
    ws.cell(row=start_row, column=1).font = HEADER_FONT
    ws.cell(row=start_row, column=1).alignment = Alignment(horizontal="center", vertical="center")

    cols_matriz = list(df_matriz.columns)
    for c_idx, col_name in enumerate(cols_matriz, start=2):
        cell = ws.cell(row=start_row, column=c_idx, value=str(col_name))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[start_row].height = 28

    for r_i, (idx_name, row_series) in enumerate(df_matriz.iterrows()):
        curr_row = start_row + 1 + r_i
        is_tot_row = (r_i == len(df_matriz) - 1 or str(idx_name).upper() == 'TOTAL')

        cell_idx = ws.cell(row=curr_row, column=1, value=str(idx_name))
        if is_tot_row:
            cell_idx.font = TOTAL_FONT
            cell_idx.fill = TOTAL_FILL
            cell_idx.border = TOTAL_BORDER
        else:
            cell_idx.font = REGULAR_FONT
            cell_idx.border = THIN_BORDER
        cell_idx.alignment = Alignment(horizontal="left", vertical="center")

        for c_i, col_name in enumerate(cols_matriz, start=2):
            val = row_series[col_name]
            cell = ws.cell(row=curr_row, column=c_i)
            cell.value = val if pd.notnull(val) else 0.0

            if is_tot_row:
                cell.font = TOTAL_FONT
                cell.fill = TOTAL_FILL
                cell.border = TOTAL_BORDER
            else:
                cell.font = REGULAR_FONT
                cell.border = THIN_BORDER

            cell.alignment = Alignment(horizontal="right", vertical="center")
            if isinstance(val, (int, float)):
                cell.number_format = '$#,##0.00'

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
    ws.column_dimensions['A'].width = 28

    wb.save(output)
    return output.getvalue()


# --- INTERFAZ DE AMARRES ---
estructura = obtener_estructura_balanzas()

if estructura:
    anios_disponibles = sorted(list(set(x['anio'] for x in estructura)), reverse=True)
    col_a, col_m = st.columns([1, 1])

    with col_a:
        anio_sel = st.selectbox("Año:", anios_disponibles)

    meses_disponibles = sorted(
        list(set(x['mes'] for x in estructura if x['anio'] == anio_sel)),
        reverse=True,
    )

    with col_m:
        mes_sel = st.selectbox("Mes:", meses_disponibles)

    ruta_balanza = next(
        (x['ruta'] for x in estructura if x['anio'] == anio_sel and x['mes'] == mes_sel),
        estructura[0]['ruta'],
    )

    st.markdown("---")

    subtab_intercos, subtab_ig_intercos = st.tabs([
        "🔄 Amarre Intercompañías",
        "📑 Amarre I y G Intercos",
    ])

    todas_empresas_balanza = obtener_lista_empresas(ruta_balanza)

    with subtab_intercos:
        st.markdown("### 🔄 Amarre Intercompañías (Cuentas de Balance: Clientes, Proveedores y Préstamos)")
        st.info(f"Cálculo estricto con saldos netos consolidados por contraparte (Deudor - Acreedor): `{ruta_balanza}`.")

        if todas_empresas_balanza:
            datos_empresas_recs = {}
            for emp in todas_empresas_balanza:
                df_b_raw = cargar_hoja_balanza(ruta_balanza, emp)
                num_cols = df_b_raw.shape[1]
                col_cta = 0
                col_deudor_f = 6 if num_cols > 6 else num_cols - 2
                col_acreedor_f = 7 if num_cols > 7 else num_cols - 1

                recs_list = []
                for _, row in df_b_raw.iterrows():
                    cta_raw = str(row.iloc[col_cta]).strip()
                    if not cta_raw or cta_raw.lower() in ('nan', 'cuenta', 'none'):
                        continue
                    if not re.search(r'\d{3}[-\s]?\d{3,5}', cta_raw):
                        continue
                    d_f = parse_monto_robusto(row.iloc[col_deudor_f])
                    a_f = parse_monto_robusto(row.iloc[col_acreedor_f])
                    
                    saldo_neto = d_f - a_f

                    recs_list.append({
                        'cta_raw': cta_raw,
                        'saldo_final': saldo_neto
                    })
                datos_empresas_recs[emp] = recs_list

            # --- MATRIZ 1: FACTURACIÓN ---
            matriz_fact = pd.DataFrame(0.0, index=todas_empresas_balanza, columns=todas_empresas_balanza)

            for emp_receptora in todas_empresas_balanza:
                recs = datos_empresas_recs[emp_receptora]
                for r in recs:
                    cta = r['cta_raw']
                    if cta.startswith(('4', '5', '6', '7')):
                        continue

                    if cta.startswith(("101-00004", "201-")) or "00004" in cta or "201-" in cta:
                        cod_contraparte = extraer_codigo_contraparte_robusto(cta)
                        emp_origen = MAPEO_CODIGO_EMPRESA.get(cod_contraparte, None)

                        if emp_origen and emp_origen in todas_empresas_balanza and emp_origen != emp_receptora:
                            monto = r['saldo_final']
                            # CORRECCIÓN: Se mantiene el saldo neto con su signo natural (Deudor - Acreedor)
                            # para que las cuentas acreedoras (Proveedores 201-) resten correctamente en lugar de sumarse.
                            matriz_fact.loc[emp_receptora, emp_origen] += monto

            matriz_fact['TOTAL'] = matriz_fact.sum(axis=1)
            matriz_fact.loc['TOTAL'] = matriz_fact.sum(axis=0)

            st.markdown("#### 📄 Matriz de Clientes y Proveedores (Facturación)")
            st.dataframe(matriz_fact.style.format('${:,.2f}'), use_container_width=True)

            excel_fact = exportar_excel_matriz_individual(
                matriz_fact,
                "REPORTE EJECUTIVO - AMARRE DE FACTURACIÓN",
                f"Periodo: {mes_sel}/{anio_sel} | Cuentas netas por empresa contraparte"
            )
            st.download_button(
                label="📥 Descargar Excel - Matriz de Facturación",
                data=excel_fact,
                file_name=f"Amarre_Facturacion_{mes_sel}_{anio_sel}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

            st.markdown("---")

            # --- MATRIZ 2: PRÉSTAMOS ---
            matriz_prest = pd.DataFrame(0.0, index=todas_empresas_balanza, columns=todas_empresas_balanza)

            for emp_receptora in todas_empresas_balanza:
                recs = datos_empresas_recs[emp_receptora]
                for r in recs:
                    cta = r['cta_raw']
                    if cta.startswith(('4', '5', '6', '7')):
                        continue

                    if cta.startswith(("101-00015", "202-00001")) or "00015" in cta or "202-00001" in cta:
                        cod_contraparte = extraer_codigo_contraparte_robusto(cta)
                        emp_origen = MAPEO_CODIGO_EMPRESA.get(cod_contraparte, None)

                        if emp_origen and emp_origen in todas_empresas_balanza and emp_origen != emp_receptora:
                            monto = r['saldo_final']
                            if cta.startswith("202-00001") or "202-" in cta:
                                monto = -abs(monto)
                            matriz_prest.loc[emp_receptora, emp_origen] += monto

            matriz_prest['TOTAL'] = matriz_prest.sum(axis=1)
            matriz_prest.loc['TOTAL'] = matriz_prest.sum(axis=0)

            st.markdown("#### 💸 Matriz de Préstamos Intercompañías")
            st.dataframe(matriz_prest.style.format('${:,.2f}'), use_container_width=True)

            excel_prest = exportar_excel_matriz_individual(
                matriz_prest,
                "REPORTE EJECUTIVO - AMARRE DE PRÉSTAMOS",
                f"Periodo: {mes_sel}/{anio_sel} | Cuentas netas por empresa contraparte"
            )
            st.download_button(
                label="📥 Descargar Excel - Matriz de Préstamos",
                data=excel_prest,
                file_name=f"Amarre_Prestamos_{mes_sel}_{anio_sel}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

    with subtab_ig_intercos:
        st.markdown("### 📑 Amarre I y G Intercos (Ingresos facturados y Costos Intercos)")
        st.info("Ingresos Facturados vs. Ingresos Contabilidad (Cuentas 410, 411, 423 menos Notas de Crédito 420, 421) y Costos Netos.")

        if todas_empresas_balanza:
            datos_empresas_recs = {}
            for emp in todas_empresas_balanza:
                df_b_raw = cargar_hoja_balanza(ruta_balanza, emp)
                num_cols = df_b_raw.shape[1]
                col_cta = 0
                col_deudor_f = 6 if num_cols > 6 else num_cols - 2
                col_acreedor_f = 7 if num_cols > 7 else num_cols - 1

                recs_list = []
                for _, row in df_b_raw.iterrows():
                    cta_raw = str(row.iloc[col_cta]).strip()
                    if not cta_raw or cta_raw.lower() in ('nan', 'cuenta', 'none'):
                        continue
                    if not re.search(r'\d{3}[-\s]?\d{3,5}', cta_raw):
                        continue
                    d_f = parse_monto_robusto(row.iloc[col_deudor_f])
                    a_f = parse_monto_robusto(row.iloc[col_acreedor_f])
                    recs_list.append({
                        'cta_raw': cta_raw,
                        'deudor_f': d_f,
                        'acreedor_f': a_f,
                        'saldo_final': d_f - a_f
                    })
                datos_empresas_recs[emp] = recs_list

            filas_ig = []
            tot_ingresos_gen = 0.0
            tot_ingresos_cont_gen = 0.0
            tot_costos_base_gen = 0.0
            tot_cta_054_gen = 0.0
            tot_costo_neto_gen = 0.0
            tot_diferencia_gen = 0.0

            for emp_destino in todas_empresas_balanza:
                cod_destino = MAPEO_NOMBRE_A_CODIGO.get(emp_destino, "")

                ingresos_facturados_a_emp = 0.0
                if cod_destino:
                    for emp_facturadora in todas_empresas_balanza:
                        if emp_facturadora == emp_destino:
                            continue
                        recs_f = datos_empresas_recs[emp_facturadora]
                        for r in recs_f:
                            cta = r['cta_raw'].strip()
                            if cta.startswith('4'):
                                segs = cta.split('-')
                                if any(seg.strip() == cod_destino for seg in segs):
                                    ingresos_facturados_a_emp += (r.get('acreedor_f', 0.0) - r.get('deudor_f', 0.0))

                ingresos_contables_emp = 0.0
                cuentas_ing_proc = set()

                costos_base_emp = 0.0
                cta_054_emp = 0.0
                cuentas_base_proc = set()
                cuentas_054_proc = set()
                
                recs_d = datos_empresas_recs[emp_destino]
                for r in recs_d:
                    cta = r['cta_raw'].strip()
                    segs = cta.split('-')
                    
                    if len(segs) >= 4:
                        prefix = segs[0].strip()
                        seg2 = segs[1].strip()
                        seg3 = segs[2].strip()
                        seg4 = segs[3].strip()

                        if prefix.isdigit() and len(prefix) == 3:
                            es_base = (seg2 in ('00000', '0000', '0') and seg3 in ('000', '0') and seg4 in ('0000', '0'))
                            saldo_cta = r.get('acreedor_f', 0.0) - r.get('deudor_f', 0.0)

                            if prefix in ('410', '411', '423') and es_base and cta not in cuentas_ing_proc:
                                ingresos_contables_emp += saldo_cta
                                cuentas_ing_proc.add(cta)
                            elif prefix in ('420', '421') and es_base and cta not in cuentas_ing_proc:
                                ingresos_contables_emp -= saldo_cta
                                cuentas_ing_proc.add(cta)
                            elif prefix.startswith(('5', '6')) and prefix.endswith('1'):
                                es_cuenta_base = es_base
                                es_054_exacta = (seg3 in ('054', '54', '0054') and seg4 in ('0000', '0'))

                                if es_cuenta_base and cta not in cuentas_base_proc:
                                    costos_base_emp += r['saldo_final']
                                    cuentas_base_proc.add(cta)
                                elif es_054_exacta and cta not in cuentas_054_proc:
                                    cta_054_emp += r['saldo_final']
                                    cuentas_054_proc.add(cta)

                costo_neto_emp = costos_base_emp - cta_054_emp
                diferencia_emp = ingresos_facturados_a_emp - costo_neto_emp

                tot_ingresos_gen += ingresos_facturados_a_emp
                tot_ingresos_cont_gen += ingresos_contables_emp
                tot_costos_base_gen += costos_base_emp
                tot_cta_054_gen += cta_054_emp
                tot_costo_neto_gen += costo_neto_emp
                tot_diferencia_gen += diferencia_emp

                filas_ig.append({
                    'EMPRESA': emp_destino,
                    'INGRESOS FACTURADOS': ingresos_facturados_a_emp,
                    'INGRESOS CONTABILIDAD': ingresos_contables_emp,
                    'COSTOS BASE (-00000-)': costos_base_emp,
                    'CUENTA 054': cta_054_emp,
                    'COSTO NETO (COSTOS - 054)': costo_neto_emp,
                    'DIFERENCIA (INGRESOS - COSTO NETO)': diferencia_emp,
                })

            df_ig_resumen = pd.DataFrame(filas_ig)
            
            df_ig_resumen.loc[len(df_ig_resumen)] = {
                'EMPRESA': 'TOTAL',
                'INGRESOS FACTURADOS': tot_ingresos_gen,
                'INGRESOS CONTABILIDAD': tot_ingresos_cont_gen,
                'COSTOS BASE (-00000-)': tot_costos_base_gen,
                'CUENTA 054': tot_cta_054_gen,
                'COSTO NETO (COSTOS - 054)': tot_costo_neto_gen,
                'DIFERENCIA (INGRESOS - COSTO NETO)': tot_diferencia_gen,
            }

            st.dataframe(
                df_ig_resumen.style.format({
                    'INGRESOS FACTURADOS': '${:,.2f}',                     'INGRESOS CONTABILIDAD': '${:,.2f}',
                    'COSTOS BASE (-00000-)': '${:,.2f}',                     'CUENTA 054': '${:,.2f}',
                    'COSTO NETO (COSTOS - 054)': '${:,.2f}',                     'DIFERENCIA (INGRESOS - COSTO NETO)': '${:,.2f}',
                }),
                use_container_width=True,
                hide_index=True
            )

            excel_ig = exportar_excel_matriz_individual(
                df_ig_resumen.set_index('EMPRESA'),
                "REPORTE EJECUTIVO - AMARRE INGRESOS VS COSTOS/GASTOS (I y G)",
                f"Periodo: {mes_sel}/{anio_sel} | Resumen consolidado por empresa"
            )
            st.download_button(
                label="📥 Descargar Excel - Amarre I y G Intercos",
                data=excel_ig,
                file_name=f"Amarre_IyG_Intercos_{mes_sel}_{anio_sel}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )

            st.success("✅ Cuentas de ingresos, notas de crédito y costos intercos barridas y calculadas correctamente.")
else:
    st.info("Por favor coloca tus carpetas 'Balanzas' y el archivo 'FORMATO EDO RESULTADOS.xlsx' en la raíz.")