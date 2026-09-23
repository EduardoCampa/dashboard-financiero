import streamlit as st
import pandas as pd
import io
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Reporte Ejecutivo de Pagos - Grupo SERVYRE", layout="wide")

def formato_mx(val):
    if pd.isnull(val): return "$0.00"
    try:
        num = float(val)
        partes = f"{num:,.2f}".split(".")
        return f"${int(partes[0].replace(',', '')):,}.{partes[1] if len(partes) > 1 else '00'}"
    except (ValueError, TypeError):
        return str(val)

def generar_excel_ejecutivo(df_datos, empresa_nombre):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Pagos"
    ws.views.sheetView[0].showGridLines = True

    font_titulo = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    fill_titulo = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    font_empresa = Font(name="Calibri", size=12, bold=True, color="1F497D")
    fill_empresa = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    font_prov = Font(name="Calibri", size=11, bold=True, color="333333")
    fill_prov = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    font_moneda = Font(name="Calibri", size=10, bold=True, color="1F497D")
    fill_moneda = PatternFill(start_color="E9EDF4", end_color="E9EDF4", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    font_normal = Font(name="Calibri", size=10, color="000000")
    font_total = Font(name="Calibri", size=11, bold=True, color="000000")
    fill_total = PatternFill(start_color="8EA9DB", end_color="8EA9DB", fill_type="solid")

    borde_delgado = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'), top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))
    borde_total = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))

    row_idx = 1
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=8)
    cell = ws.cell(row=row_idx, column=1, value=f"REPORTE EJECUTIVO DE PAGOS — {empresa_nombre.upper()}")
    cell.font = font_titulo; cell.fill = fill_titulo; cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row_idx].height = 35
    row_idx += 2

    headers = ["TIPO", "PROVEEDOR", "FECHA VENCIMIENTO", "FOLIO / DOCUMENTO", "UUID", "MONEDA", "DESCRIPCIÓN", "SALDO PENDIENTE"]
    col_emp = 'EmpresaOrigen' if 'EmpresaOrigen' in df_datos.columns else None
    col_prov = 'BusinessEntityName' if 'BusinessEntityName' in df_datos.columns else None
    col_curr = 'Currency' if 'Currency' in df_datos.columns else None

    empresas = sorted(df_datos[col_emp].unique()) if col_emp else ['General']

    for empresa in empresas:
        df_emp = df_datos[df_datos[col_emp] == empresa] if col_emp else df_datos
        total_empresa = df_emp['Saldo_Pendiente'].sum()

        ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
        cell_emp = ws.cell(row=row_idx, column=1, value=f"EMPRESA: {str(empresa).upper()}")
        cell_emp.font = font_empresa; cell_emp.fill = fill_empresa; cell_emp.alignment = Alignment(horizontal="left", vertical="center")
        
        cell_emp_tot = ws.cell(row=row_idx, column=8, value=total_empresa)
        cell_emp_tot.font = font_empresa; cell_emp_tot.fill = fill_empresa; cell_emp_tot.number_format = '"$"#,##0.00'; cell_emp_tot.alignment = Alignment(horizontal="right", vertical="center")
        
        for c_idx in range(1, 9): ws.cell(row=row_idx, column=c_idx).border = borde_delgado
        ws.row_dimensions[row_idx].height = 26
        row_idx += 1

        proveedores = sorted(df_emp[col_prov].unique()) if col_prov else []
        for prov in proveedores:
            df_prov = df_emp[df_emp[col_prov] == prov]
            total_prov = df_prov['Saldo_Pendiente'].sum()

            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
            cell_prov = ws.cell(row=row_idx, column=1, value=f"   👤 {str(prov)}")
            cell_prov.font = font_prov; cell_prov.fill = fill_prov; cell_prov.alignment = Alignment(horizontal="left", vertical="center")
            
            cell_prov_tot = ws.cell(row=row_idx, column=8, value=total_prov)
            cell_prov_tot.font = font_prov; cell_prov_tot.fill = fill_prov; cell_prov_tot.number_format = '"$"#,##0.00'; cell_prov_tot.alignment = Alignment(horizontal="right", vertical="center")
            
            for c_idx in range(1, 9): ws.cell(row=row_idx, column=c_idx).border = borde_delgado
            ws.row_dimensions[row_idx].height = 22
            row_idx += 1

            monedas_prov = sorted(df_prov[col_curr].dropna().unique()) if col_curr else ['MXN']
            for moneda in monedas_prov:
                df_moneda = df_prov[df_prov[col_curr] == moneda] if col_curr else df_prov
                total_moneda = df_moneda['Saldo_Pendiente'].sum()

                ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=7)
                cell_mon = ws.cell(row=row_idx, column=1, value=f"      💱 Moneda: {str(moneda)}")
                cell_mon.font = font_moneda; cell_mon.fill = fill_moneda; cell_mon.alignment = Alignment(horizontal="left", vertical="center")
                
                cell_mon_tot = ws.cell(row=row_idx, column=8, value=total_moneda)
                cell_mon_tot.font = font_moneda; cell_mon_tot.fill = fill_moneda; cell_mon_tot.number_format = '"$"#,##0.00'; cell_mon_tot.alignment = Alignment(horizontal="right", vertical="center")
                
                for c_idx in range(1, 9): ws.cell(row=row_idx, column=c_idx).border = borde_delgado
                ws.row_dimensions[row_idx].height = 20
                row_idx += 1

                for col_num, h in enumerate(headers, 1):
                    c = ws.cell(row=row_idx, column=col_num, value=h)
                    c.font = font_header; c.fill = fill_header; c.alignment = Alignment(horizontal="center", vertical="center"); c.border = borde_delgado
                ws.row_dimensions[row_idx].height = 20
                row_idx += 1

                for _, r in df_moneda.iterrows():
                    ws.cell(row=row_idx, column=1, value=str(r.get('Tipo_Movimiento', ''))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=2, value=str(r.get('BusinessEntityName', ''))).alignment = Alignment(horizontal="left")
                    ws.cell(row=row_idx, column=3, value=str(r.get('Fecha_Fmt', ''))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=4, value=str(r.get('DocFolio', ''))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=5, value=str(r.get('UUID', ''))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=6, value=str(r.get('Currency', 'MXN'))).alignment = Alignment(horizontal="center")
                    ws.cell(row=row_idx, column=7, value=str(r.get('Title', ''))).alignment = Alignment(horizontal="left")
                    
                    c_val = ws.cell(row=row_idx, column=8, value=float(r.get('Saldo_Pendiente', 0)))
                    c_val.number_format = '"$"#,##0.00'; c_val.alignment = Alignment(horizontal="right")

                    for col_num in range(1, 9): ws.cell(row=row_idx, column=col_num).font = font_normal; ws.cell(row=row_idx, column=col_num).border = borde_delgado
                    ws.row_dimensions[row_idx].height = 18
                    row_idx += 1
                row_idx += 1
            row_idx += 1
        row_idx += 1

    ws.cell(row=row_idx, column=1, value="TOTAL GENERAL").font = font_total
    for col_num in range(1, 8): ws.cell(row=row_idx, column=col_num).fill = fill_total; ws.cell(row=row_idx, column=col_num).border = borde_total
    
    cell_tot_gen = ws.cell(row=row_idx, column=8, value=df_datos['Saldo_Pendiente'].sum())
    cell_tot_gen.font = font_total; cell_tot_gen.fill = fill_total; cell_tot_gen.number_format = '"$"#,##0.00'
    cell_tot_gen.border = borde_total; cell_tot_gen.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[row_idx].height = 26

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

ruta_archivo = "../Consolidado_Master.xlsx" if not os.path.exists("Consolidado_Master.xlsx") else "Consolidado_Master.xlsx"

@st.cache_data
def cargar_reporte():
    if not os.path.exists(ruta_archivo): return None, None, None, None, None
    try:
        xls = pd.ExcelFile(ruta_archivo)
        sheets = xls.sheet_names
        return (
            pd.read_excel(ruta_archivo, sheet_name='SolicitudPago') if 'SolicitudPago' in sheets else pd.DataFrame(),
            pd.read_excel(ruta_archivo, sheet_name='OrdenCompra') if 'OrdenCompra' in sheets else pd.DataFrame(),
            pd.read_excel(ruta_archivo, sheet_name='EdoCuenta') if 'EdoCuenta' in sheets else pd.DataFrame(),
            pd.read_excel(ruta_archivo, sheet_name='FacturaCompra') if 'FacturaCompra' in sheets else None,
            pd.read_excel(ruta_archivo, sheet_name='Gastos') if 'Gastos' in sheets else None
        )
    except Exception:
        return None, None, None, None, None

df_tesoreria, df_ordenes, df_edocuenta, df_fact_compra, df_gastos = cargar_reporte()

st.title("📋 Reporte Ejecutivo de Pagos con Saldo Pendiente")

# Procesamiento unificado para el reporte
df_oc_base = pd.DataFrame()
if df_ordenes is not None and not df_ordenes.empty:
    df_oc_base = df_ordenes.copy()
    df_oc_base['Tipo_Movimiento'] = 'Orden de Compra (OC)'
    df_oc_base['Saldo_Pendiente'] = pd.to_numeric(df_oc_base.get('Saldo_Pendiente', 0), errors='coerce').fillna(0)

df_sp_base = pd.DataFrame()
if df_tesoreria is not None and not df_tesoreria.empty:
    df_sp_base = df_tesoreria.copy()
    df_sp_base['Tipo_Movimiento'] = 'Solicitud de Pago (SP)'
    df_sp_base['Saldo_Pendiente'] = pd.to_numeric(df_sp_base.get('Saldo_Pendiente', 0), errors='coerce').fillna(0)

df_rep_list = [df for df in [df_oc_base, df_sp_base] if not df.empty]

if df_rep_list:
    df_rep_total = pd.concat(df_rep_list, ignore_index=True)
    df_rep_total = df_rep_total[df_rep_total['Saldo_Pendiente'] > 1.0]

    col_emp = 'EmpresaOrigen' if 'EmpresaOrigen' in df_rep_total.columns else None
    col_prov = 'BusinessEntityName' if 'BusinessEntityName' in df_rep_total.columns else None
    col_folio = 'DocFolio' if 'DocFolio' in df_rep_total.columns else None

    if col_prov and col_folio:
        st.markdown("#### ⚙️ Filtros Independientes")
        c1, c2, c3 = st.columns(3)
        with c1:
            if col_emp:
                e_sel = st.multiselect("Empresa:", sorted(df_rep_total[col_emp].dropna().unique()), key="rep_e")
                if e_sel: df_rep_total = df_rep_total[df_rep_total[col_emp].isin(e_sel)]
        with c2:
            p_sel = st.multiselect("Proveedor:", sorted(df_rep_total[col_prov].dropna().unique()), key="rep_p")
            if p_sel: df_rep_total = df_rep_total[df_rep_total[col_prov].isin(p_sel)]
        with c3:
            t_sel = st.multiselect("Tipo:", sorted(df_rep_total['Tipo_Movimiento'].dropna().unique()), key="rep_t")
            if t_sel: df_rep_total = df_rep_total[df_rep_total['Tipo_Movimiento'].isin(t_sel)]

        st.markdown("---")
        st.markdown(f"### 💰 **Total Saldo Pendiente General:** {formato_mx(df_rep_total['Saldo_Pendiente'].sum())}")
        
        empresa_para_excel = df_rep_total[col_emp].iloc[0] if col_emp and len(df_rep_total[col_emp].dropna().unique()) == 1 else "Consolidado"
        st.download_button(
            label="📥 Descargar Reporte en Excel con UUID",
            data=generar_excel_ejecutivo(df_rep_total, empresa_para_excel),
            file_name="Reporte_Ejecutivo_Pagos_UUID.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        st.markdown("---")
        st.dataframe(df_rep_total, use_container_width=True)
else:
    st.warning("No hay registros pendientes para el reporte.")